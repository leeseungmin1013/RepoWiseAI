from pathlib import Path
from datetime import date
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "project-status"
SHOT_DIR = OUT_DIR / "screenshots"
OUT_FILE = OUT_DIR / "RepoWiseAI_프로젝트_진행_상황_보고서.docx"

FONT = "Malgun Gothic"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5B6573"
LIGHT = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "2F6B4F"
AMBER = "8A6100"
RED = "9B1C1C"
USABLE_DXA = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == USABLE_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag in ("tblW", "tblInd", "tblLayout"):
        old = tbl_pr.find(qn(f"w:{tag}"))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color=None, bold=None):
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, 11, "222222")
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        set_style_font(style, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.styles["Title"]
    set_style_font(title, 28, NAVY, True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(7)

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, 13, MUTED, False)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)

    caption = doc.styles["Caption"]
    set_style_font(caption, 9, MUTED, False)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_numbering(doc, bullet=False):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abs_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "●" if bullet else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    if bullet:
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        fonts.set(qn("w:eastAsia"), FONT)
        fonts.set(qn("w:hint"), "default")
        r_pr.append(fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def patch_existing_numbering(doc, num_id):
    numbering = doc.part.numbering_part.element
    num = next(x for x in numbering.findall(qn("w:num")) if x.get(qn("w:numId")) == str(num_id))
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(x for x in numbering.findall(qn("w:abstractNum")) if x.get(qn("w:abstractNumId")) == abstract_id)
    lvl = abstract.find(qn("w:lvl"))
    p_pr = lvl.find(qn("w:pPr"))
    ind = p_pr.find(qn("w:ind"))
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    tabs = p_pr.find(qn("w:tabs"))
    tab = tabs.find(qn("w:tab"))
    tab.set(qn("w:pos"), "720")
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def add_bullet(doc, text, num_id, bold_lead=None):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True, color=NAVY)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_callout(doc, label, text, fill=CALLOUT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [USABLE_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + "  ")
    set_run_font(r, 10.5, True, accent)
    r = p.add_run(text)
    set_run_font(r, 10.5, False, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_status_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 2100, 5360])
    headers = ["구분", "현재 상태", "설명"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, 9.5, True, NAVY)
    set_repeat_table_header(table.rows[0])
    for label, status, desc in rows:
        cells = table.add_row().cells
        for idx, text in enumerate((label, status, desc)):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            color = GREEN if status == "완료" else AMBER if status == "부분 완료" else RED if status == "미구현" else "222222"
            r = p.add_run(text)
            set_run_font(r, 9.2, idx == 1, color if idx == 1 else "222222")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc, filename, caption, width=6.35):
    path = SHOT_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    extent = run._element.xpath(".//wp:docPr")
    if extent:
        extent[0].set("descr", caption)
    c = doc.add_paragraph(caption, style="Caption")
    return c


def add_page_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.page_break_before = True
    return p


def set_headers_footers(doc):
    even_odd = doc.settings._element.find(qn("w:evenAndOddHeaders"))
    if even_odd is not None:
        doc.settings._element.remove(even_odd)
    for section in doc.sections:
        section.different_first_page_header_footer = False
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        hp = section.header.paragraphs[0]
        hp.text = "RepoWise AI  |  프로젝트 진행 상황 보고서"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.paragraph_format.space_after = Pt(0)
        for run in hp.runs:
            set_run_font(run, 8.5, False, MUTED)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = fp.add_run("페이지 ")
        set_run_font(r, 8.5, False, MUTED)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r2 = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r2.append(t)
        fld.append(r2)
        fp._p.append(fld)


def add_title_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("PROJECT STATUS REPORT")
    set_run_font(r, 10, True, BLUE)
    title = doc.add_paragraph("RepoWise AI\n프로젝트 진행 상황 보고서", style="Title")
    title.paragraph_format.line_spacing = 1.05
    sub = doc.add_paragraph("GitHub 저장소를 이해하고 학습하는 AI 기반 코드 탐색 서비스", style="Subtitle")
    sub.paragraph_format.space_after = Pt(22)

    add_callout(
        doc,
        "현재 단계",
        "로컬 개발 환경에서 핵심 기능 흐름 구현 완료. 현재는 품질·복구·운영 준비를 강화하는 단계입니다.",
    )

    meta = [
        ("기준일", "2026년 8월 4일"),
        ("개발 형태", "1인 기획·개발 / Codex 등 AI를 전적으로 활용"),
        ("분석 대상", "github.com/leeseungmin1013/RepoWiseAI"),
        ("문서 대상", "프로젝트를 처음 접하는 비전공자"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{label}: ")
        set_run_font(r, 10.5, True, NAVY)
        r = p.add_run(value)
        set_run_font(r, 10.5, False, "333333")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("작성 근거: docs/PROJECT_OVERVIEW.md(2026-08-02) 및 실제 RepoWiseAI 분석 화면")
    set_run_font(r, 9, False, MUTED)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    bullet_id = add_numbering(doc, bullet=True)
    decimal_id = add_numbering(doc, bullet=False)
    set_headers_footers(doc)
    add_title_page(doc)

    add_page_heading(doc, "1. 프로젝트 동기와 목적", level=1)
    doc.add_paragraph(
        "GitHub 저장소에는 서비스의 실제 동작을 설명하는 코드가 들어 있지만, 처음 보는 사람은 파일명과 기술 용어만으로 전체 구조를 이해하기 어렵습니다. RepoWise AI는 저장소의 목적과 기능 흐름을 먼저 보여 주고, 필요한 경우 실제 코드까지 단계적으로 확인할 수 있도록 만든 프로젝트입니다."
    )
    add_bullet(doc, "저장소의 목적과 주요 역할을 쉬운 문장과 구조도로 설명합니다.", bullet_id)
    add_bullet(doc, "기능이 어떤 순서로 동작하는지 정상 경로와 실패 경로로 나누어 보여 줍니다.", bullet_id)
    add_bullet(doc, "질문과 설명은 현재 분석한 코드 위치를 근거로 연결합니다.", bullet_id)
    add_bullet(doc, "사용자 수준을 진단하고 필요한 개념부터 학습 순서를 조정합니다.", bullet_id)

    doc.add_heading("AI 기반 바이브 코딩에 대한 도전", level=2)
    doc.add_paragraph(
        "기능 구현 자체와 함께, 한 명의 개발자가 Codex 등 AI를 적극적으로 활용했을 때 기획부터 화면·서버·데이터·테스트·문서화까지 어느 수준의 제품을 만들 수 있는지 확인하는 것도 중요한 동기였습니다. AI는 코드 탐색, 초안 작성, 수정, 테스트와 문서화를 지원했고, 기능 범위와 우선순위 결정 및 최종 확인은 개발자가 담당했습니다."
    )

    doc.add_heading("2. 참여 인원과 역할", level=1)
    add_status_table(doc, [
        ("참여 인원", "1명", "기획, 설계, 개발, 검증, 문서화를 모두 담당"),
        ("개발자", "총괄", "요구사항·우선순위 결정, 화면과 기능 설계, 결과 검토"),
        ("Codex 등 AI", "개발 지원", "코드 분석·작성·수정, 테스트, 문서 정리 지원"),
    ])
    add_callout(doc, "역할 구분", "AI는 개발을 돕는 도구이며, 프로젝트의 방향과 최종 결과에 대한 판단과 책임은 개발자에게 있습니다.")

    doc.add_heading("3. 현재 상태 요약", level=1)
    add_status_table(doc, [
        ("저장소 분석", "완료", "공개 GitHub URL 등록, 고정 버전 분석, 파일·기호·관계 저장"),
        ("구조·기능 탐색", "완료", "목적, 역할, 구현 구조, 기능 흐름, 코드 이동"),
        ("질문·변경 영향", "완료", "코드 근거 답변, 변경 영향 범위, 장시간 작업 연결"),
        ("적응형 학습", "부분 완료", "진단·학습 경로·활동은 구현, 난이도 조정과 유형 확대 필요"),
        ("음성·공식 자료", "부분 완료", "기본 연결은 구현, 세션 저장과 최신성 표시 필요"),
        ("인증·배포", "미구현", "로그인, 개인 데이터 구분, 비공개 저장소, 실제 서비스 배포 필요"),
    ])

    doc.add_heading("4. 현재까지 구현한 기능", level=1)
    doc.add_paragraph("다음 장부터 실제 RepoWiseAI 분석 화면을 기준으로 사용자에게 보이는 기능을 순서대로 설명합니다.")
    add_page_heading(doc, "4.1 RepoWiseAI URL 등록과 분석 진행", level=2)
    doc.add_paragraph(
        "사용자가 GitHub 저장소 주소를 입력하면 시스템이 저장소를 등록하고 분석 작업을 시작합니다. 이번 문서의 화면은 모두 실제 RepoWiseAI 주소를 입력해 새로 분석한 결과입니다."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run()
    r1.add_picture(str(SHOT_DIR / "00_repowise_analysis_input.png"), width=Inches(4.25))
    r2 = p.add_run("  ")
    set_run_font(r2, 8)
    r3 = p.add_run()
    r3.add_picture(str(SHOT_DIR / "00b_repowise_analysis_progress.png"), width=Inches(2.0))
    for dp, desc in ((r1, "RepoWiseAI GitHub URL 입력 화면"), (r3, "RepoWiseAI 분석 진행 화면")):
        props = dp._element.xpath(".//wp:docPr")
        if props:
            props[0].set("descr", desc)
    doc.add_paragraph("그림 1. 왼쪽: RepoWiseAI URL 입력 / 오른쪽: 분석 단계 진행", style="Caption")
    add_bullet(doc, "입력 주소: https://github.com/leeseungmin1013/RepoWiseAI", bullet_id, "입력 주소:")
    add_bullet(doc, "분석 대상 버전을 고정하여 서로 다른 코드 버전의 정보가 섞이지 않도록 합니다.", bullet_id)
    add_bullet(doc, "수집, 코드 정리, 관계 분석, 설명 생성 등의 진행 상태를 화면에서 확인할 수 있습니다.", bullet_id)
    add_callout(doc, "실제 분석 결과", "182개 파일, 1,303개 코드 요소, 1,771개 검색 단위가 확인되었습니다.")

    add_page_heading(doc, "4.2 저장소 소개와 전체 구조", level=2)
    doc.add_paragraph(
        "분석이 끝나면 저장소의 목적을 쉬운 문장으로 설명하고, 사용자와 시스템의 주요 역할을 구조도로 보여 줍니다. 비전공자는 파일 목록을 먼저 읽지 않아도 서비스가 무엇을 하는지 파악할 수 있습니다."
    )
    add_figure(doc, "01_repository_structure.png", "그림 2. RepoWiseAI 저장소 소개와 역할 중심 구조", width=5.6)
    add_bullet(doc, "저장소의 목적과 주요 기능을 화면 상단에서 요약합니다.", bullet_id)
    add_bullet(doc, "파일 수, 코드 요소 수, 검색 단위 수로 분석 규모를 확인하고, 역할을 선택해 관련 구현과 기능 흐름으로 이어서 탐색할 수 있습니다.", bullet_id)

    add_page_heading(doc, "4.3 역할 그래프와 상세 설명", level=2)
    doc.add_paragraph(
        "역할 그래프는 파일명 대신 사용자가 이해하기 쉬운 역할을 중심으로 저장소를 나눕니다. 화면에서는 ‘사용자 작업 공간’을 선택해 해당 역할의 설명과 관련 기능을 확인한 상태입니다."
    )
    add_figure(doc, "02_role_detail.png", "그림 3. 역할 선택 후 상세 설명을 확인하는 화면", width=5.6)
    add_bullet(doc, "사용자, 웹 화면, API, 작업 처리, 데이터 저장 등 역할 단위로 구조를 구분합니다.", bullet_id)
    add_bullet(doc, "선택한 역할의 책임과 관련 기능을 오른쪽 패널에서 설명합니다.", bullet_id)
    add_bullet(doc, "필요할 때만 실제 파일과 구현 세부 사항으로 이동할 수 있습니다.", bullet_id)

    add_page_heading(doc, "4.4 구현 구조 그래프", level=2)
    doc.add_paragraph(
        "구현 구조 그래프는 화면, API, 작업 처리, 데이터 저장 등 기술 요소가 어떻게 연결되는지 보여 줍니다. 전체 보기와 역할 중심 보기를 전환하고, 이전 분석 결과와 차이를 비교하거나 이미지로 내보낼 수 있습니다."
    )
    add_figure(doc, "03_implementation_graph.png", "그림 4. RepoWiseAI 전체 구현 구조 그래프")
    add_bullet(doc, "22개 주요 구성 요소와 48개 연결 관계를 한 화면에서 확인할 수 있습니다.", bullet_id)
    add_bullet(doc, "구성 요소를 누르면 관련 설명과 근거 코드로 이동합니다.", bullet_id)
    add_bullet(doc, "PNG 또는 Mermaid 형식으로 내보내 문서와 협업 자료에 활용할 수 있습니다.", bullet_id)

    add_page_heading(doc, "4.5 기능 흐름 확인", level=2)
    doc.add_paragraph(
        "기능 흐름은 사용자의 요청이 어떤 화면과 서버 기능을 거쳐 처리되는지 순서대로 설명합니다. 화면에서는 저장소 요청 흐름을 선택해 관련 단계와 구현 위치를 확인하고 있습니다."
    )
    add_figure(doc, "07_feature_flow.png", "그림 5. 저장소 요청 기능 흐름을 선택한 화면")
    add_bullet(doc, "기능별 시작점, 처리 단계, 결과를 하나의 흐름으로 연결합니다.", bullet_id)
    add_bullet(doc, "정상 처리뿐 아니라 오류가 발생할 수 있는 경로도 구분해 보여 줍니다.", bullet_id)
    add_bullet(doc, "각 단계에서 실제 코드 위치로 바로 이동할 수 있습니다.", bullet_id)

    add_page_heading(doc, "4.6 코드 탐색과 시작 지점 안내", level=2)
    doc.add_paragraph(
        "코드 탐색 화면은 폴더 구조, 코드 내용, 선택한 코드의 설명을 함께 제공합니다. ‘Start Here’ 안내는 처음 읽을 파일과 순서를 제시해 큰 저장소에서도 탐색을 시작하기 쉽게 합니다."
    )
    add_figure(doc, "04_code_explorer.png", "그림 6. RepoWiseAI 코드 탐색기와 Start Here 안내")
    add_bullet(doc, "왼쪽에서 파일을 선택하고 가운데에서 실제 코드를 확인합니다.", bullet_id)
    add_bullet(doc, "선택한 기능과 관련된 코드 범위를 강조해 불필요한 탐색을 줄입니다.", bullet_id)
    add_bullet(doc, "코드 일부를 선택해 해당 부분을 기준으로 질문할 수 있습니다.", bullet_id)

    add_page_heading(doc, "4.7 변경 영향 분석", level=2)
    doc.add_paragraph(
        "변경 영향 분석은 특정 파일이나 기능을 수정했을 때 함께 확인해야 할 부분을 정리합니다. 개발자는 변경 목적을 입력하고 직접 영향과 가능성이 있는 영향을 구분해 검토할 수 있습니다."
    )
    add_figure(doc, "05_change_impact.png", "그림 7. 코드 선택과 변경 영향 분석 입력 화면")
    add_bullet(doc, "현재 선택한 코드와 연결된 기능 및 파일을 추적합니다.", bullet_id)
    add_bullet(doc, "즉시 영향을 받는 부분과 추가 확인이 필요한 부분을 나누어 제시합니다.", bullet_id)
    add_bullet(doc, "분석 시간이 긴 작업은 DeepTask로 실행하고 진행 상태와 취소를 지원합니다.", bullet_id)

    add_page_heading(doc, "4.8 사용자 수준 진단과 학습 경로", level=2)
    doc.add_paragraph(
        "학습 기능은 사용자의 경험과 목표를 간단히 진단한 뒤, 현재 저장소를 이해하는 데 필요한 개념을 순서대로 구성합니다. 학습 활동 결과에 따라 보충 설명이나 다음 단계가 조정됩니다."
    )
    add_figure(doc, "06_learning_diagnostic.png", "그림 8. RepoWiseAI 적응형 학습 진단 화면")
    add_bullet(doc, "6~7개 질문으로 현재 경험 수준과 학습 목적을 확인합니다.", bullet_id)
    add_bullet(doc, "저장소의 실제 코드에 연결된 최대 40단계 학습 경로를 구성합니다.", bullet_id)
    add_bullet(doc, "틀린 답이나 부족한 선수 개념을 발견하면 보충 학습 후 원래 경로로 돌아옵니다.", bullet_id)
    add_bullet(doc, "현재는 기본 학습 활동이 구현되어 있으며, 난이도 자동 조정과 활동 유형 확대가 남아 있습니다.", bullet_id)

    add_page_heading(doc, "4.9 화면 뒤에서 구현된 핵심 기능", level=2)
    doc.add_paragraph("화면에 직접 드러나지 않지만, 서비스의 정확성과 안정성을 위해 다음 기능도 구현되어 있습니다.")
    add_bullet(doc, "코드 근거 답변: 검색 결과의 파일, 줄 번호와 내용이 현재 분석 버전과 일치하는지 확인합니다.", bullet_id, "코드 근거 답변:")
    add_bullet(doc, "고정 버전 분석: 브랜치 이름이 아니라 실제 커밋 번호를 기준으로 결과를 저장합니다.", bullet_id, "고정 버전 분석:")
    add_bullet(doc, "안전한 저장소 수집: 경로 이탈, 비밀 파일, 지나치게 큰 파일과 압축 파일을 걸러냅니다.", bullet_id, "안전한 저장소 수집:")
    add_bullet(doc, "혼합 검색: 정확한 단어 검색과 의미 기반 검색을 함께 사용합니다.", bullet_id, "혼합 검색:")
    add_bullet(doc, "장시간 작업: 데이터베이스에 작업 상태를 저장하고 진행 상황 전송, 재연결, 취소를 지원합니다.", bullet_id, "장시간 작업:")
    add_bullet(doc, "음성 학습 기반: 누르고 말하기 방식의 실시간 음성 연결을 학습 화면과 연동했습니다.", bullet_id, "음성 학습 기반:")

    doc.add_heading("기능별 현재 상태", level=2)
    add_status_table(doc, [
        ("Repository Story", "완료", "목적·역할 구조와 반응형 상세 화면"),
        ("Project Map·Feature Flow", "완료", "기능 목록, 정상·실패 경로, 코드 이동"),
        ("Code Focus·Change Brief", "완료", "최소 코드 설명과 변경 영향 추적"),
        ("Architecture Graph", "완료", "구조 검증, 비교, PNG·Mermaid 내보내기"),
        ("DeepTask", "완료", "진행 상태 전송, 취소, 재접속, 중복 방지"),
        ("실시간 음성", "부분 완료", "기본 연결 완료, 세션 저장과 명령 처리 보강 필요"),
        ("공식 학습자료", "부분 완료", "허용 목록은 있음, 최신성 자동 확인 필요"),
        ("품질 평가", "부분 완료", "구조·검색 평가는 있음, 학습·브라우저 평가 확대 필요"),
    ])

    add_page_heading(doc, "5. 남은 작업", level=1)
    add_callout(doc, "핵심 판단", "현재 버전은 로컬에서 전체 흐름을 확인할 수 있는 수준입니다. 실제 공개 서비스로 운영하려면 로그인, 배포, 보안, 복구와 운영 관측 기능이 필요합니다.", fill="FFF8E6", accent=AMBER)

    doc.add_heading("P0. 현재 기능 안정화", level=2)
    add_bullet(doc, "서로 다른 유형의 저장소에서도 구조 설명이 안정적으로 만들어지는지 검수합니다.", bullet_id)
    add_bullet(doc, "학습 화면을 새로고침해도 최근 답변, 활동과 보충 학습 상태가 복구되도록 합니다.", bullet_id)
    add_bullet(doc, "공식 학습자료의 최신 확인 시각과 사용 가능 상태를 자동으로 갱신합니다.", bullet_id)
    add_bullet(doc, "학습 품질을 변경 전후로 비교할 수 있는 평가 자료와 보고서를 만듭니다.", bullet_id)
    add_bullet(doc, "검색, AI 생성, 작업 처리 실패를 같은 기록으로 추적하고 재시도·시간 초과·재연결을 검증합니다.", bullet_id)

    doc.add_heading("P1. 계획된 기능 확장", level=2)
    add_bullet(doc, "음성 세션 저장, 말 끊기 처리, 중복 명령 방지 등 음성 기능 2단계를 완성합니다.", bullet_id)
    add_bullet(doc, "학습자의 직전 결과와 숙련도에 따라 문제 난이도와 활동 유형을 조정합니다.", bullet_id)
    add_bullet(doc, "학습 경로 변경안을 미리 보고 적용·거절하며 변경 차이를 확인하도록 합니다.", bullet_id)
    add_bullet(doc, "TypeScript와 Python 분석 범위 및 다양한 저장소 평가 자료를 확대합니다.", bullet_id)

    add_page_heading(doc, "P2. 실제 서비스 운영 준비", level=2)
    add_bullet(doc, "로그인과 인증(Authentication): 사용자 계정을 확인하고 본인의 데이터만 접근하도록 구분합니다.", bullet_id, "로그인과 인증(Authentication):")
    add_bullet(doc, "비공개 저장소: GitHub 권한 연결과 안전한 접근 범위를 구현합니다.", bullet_id, "비공개 저장소:")
    add_bullet(doc, "실제 배포: 웹, API, 데이터베이스와 작업 처리 서버를 운영 환경에 배포합니다.", bullet_id, "실제 배포:")
    add_bullet(doc, "운영 안전: 비밀 정보 관리, 데이터베이스 백업, 변경 적용과 되돌리기 절차를 마련합니다.", bullet_id, "운영 안전:")
    add_bullet(doc, "비용과 성능: AI 사용량, 비용, 처리 시간, 오류를 확인하는 화면을 만듭니다.", bullet_id, "비용과 성능:")
    add_bullet(doc, "데이터 정책: 보존 기간, 삭제, 내보내기와 자동 정리 작업을 구현합니다.", bullet_id, "데이터 정책:")
    add_bullet(doc, "보안과 품질: 요청 제한, 악용 방지, 접근성, 브라우저 호환성, 부하·보안 점검을 수행합니다.", bullet_id, "보안과 품질:")

    add_page_heading(doc, "6. 이후 진행 계획", level=1)
    doc.add_paragraph("일정은 품질 검증 결과에 따라 조정하며, 다음 순서로 진행합니다.")
    steps = [
        ("1단계", "로컬 기능 안정화", "새로고침 복구, 오류 처리, 공식 자료 최신성, 평가 자동화를 우선 보완합니다."),
        ("2단계", "로그인과 데이터 구분", "사용자 인증, 사용자별 저장소·학습 데이터 경계, 비공개 저장소 접근을 구현합니다."),
        ("3단계", "실제 배포와 운영 기반", "서버 배포, 비밀 정보 관리, 백업·변경·복구 절차, 비용·지연 관측을 준비합니다."),
        ("4단계", "기능 확장", "음성 2단계, 적응형 활동, 학습 경로 변경 제안, 분석 언어와 평가 범위를 확대합니다."),
        ("5단계", "공개 전 검증", "접근성, 주요 브라우저, 부하, 보안과 데이터 삭제 절차를 확인합니다."),
    ]
    for label, title, desc in steps:
        p = doc.add_paragraph()
        apply_numbering(p, decimal_id)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.167
        r = p.add_run(f"{title} — ")
        set_run_font(r, 11, True, NAVY)
        r = p.add_run(desc)
        set_run_font(r, 11, False, "222222")

    doc.add_heading("7. 현재 평가", level=1)
    doc.add_paragraph(
        "RepoWise AI는 공개 GitHub 저장소를 분석하고, 구조 이해부터 코드 탐색, 변경 영향 확인, 맞춤형 학습까지 연결하는 핵심 흐름을 로컬 환경에서 구현했습니다. 실제 서비스 공개 전에는 인증과 사용자 데이터 구분, 운영 환경 배포, 백업·보안·비용 관측을 반드시 보완해야 합니다."
    )
    add_callout(doc, "현재 완료 기준", "기능 시연이 가능한 로컬 제품 단계이며, 운영 준비가 끝난 상용 서비스 단계는 아닙니다.")

    add_page_heading(doc, "비전공자를 위한 용어 안내", level=2)
    add_status_table(doc, [
        ("저장소", "Repository", "프로젝트의 코드와 문서를 모아 두는 공간"),
        ("커밋", "Commit", "특정 시점에 저장된 코드 버전"),
        ("배포", "Deployment", "개발한 서비스를 인터넷에서 실제로 사용할 수 있게 올리는 과정"),
        ("인증", "Authentication", "로그인한 사용자가 누구인지 확인하는 과정"),
        ("API", "연결 창구", "웹 화면과 서버가 정보를 주고받는 규칙"),
    ])

    core = doc.core_properties
    core.title = "RepoWise AI 프로젝트 진행 상황 보고서"
    core.subject = "프로젝트 동기, 참여 역할, 구현 기능, 남은 작업 및 진행 계획"
    core.author = "RepoWiseAI 프로젝트"
    core.keywords = "RepoWiseAI, 프로젝트 진행 상황, AI, Codex, GitHub 분석"
    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    main()
